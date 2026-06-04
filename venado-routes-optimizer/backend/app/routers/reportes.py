import csv
import io
from datetime import date, datetime
from enum import Enum
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from geoalchemy2.shape import to_shape
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import Session, joinedload

from app.auth import require_role
from app.database import get_db
from app.models import EjecucionMicroTarea, ReporteIncidencia, Ruta, Visita, VisitaEstado


router = APIRouter(prefix="/reportes", tags=["reportes"])
LOCAL_TZ = ZoneInfo("America/La_Paz")
DIAS_ES = ["lunes", "martes", "miercoles", "jueves", "viernes", "sabado", "domingo"]


class TablaBI(str, Enum):
    visitas = "visitas"
    micro_tareas = "micro_tareas"
    rutas = "rutas"


VISITAS_HEADERS = [
    "fecha",
    "anio",
    "mes",
    "semana_iso",
    "dia_semana",
    "orden_visita",
    "estado_visita",
    "supervisor",
    "reponedor",
    "pdv_codigo",
    "pdv_nombre",
    "mercado",
    "tipo_cliente",
    "latitud_pdv",
    "longitud_pdv",
    "checkin_latitud",
    "checkin_longitud",
    "checkout_latitud",
    "checkout_longitud",
    "distancia_planificada_km",
    "distancia_ors_km",
    "tiempo_traslado_planificado_min",
    "tiempo_traslado_ors_min",
    "tiempo_traslado_real_min",
    "fuente_calculo_traslado",
    "hora_inicio_traslado",
    "hora_fin_traslado",
    "hora_inicio_visita",
    "hora_fin_visita",
    "tiempo_pdv_estimado_min",
    "tiempo_pdv_real_min",
    "desviacion_pdv_min",
    "desviacion_pdv_pct",
    "desviacion_pdv_alerta",
    "micro_tareas_total",
    "micro_tareas_completadas",
    "micro_tareas_pendientes",
    "evidencias_fotos_total",
    "evidencias_fotos_urls",
]

MICRO_TAREAS_HEADERS = [
    "fecha",
    "anio",
    "mes",
    "semana_iso",
    "dia_semana",
    "orden_visita",
    "estado_visita",
    "supervisor",
    "reponedor",
    "pdv_codigo",
    "pdv_nombre",
    "mercado",
    "tipo_cliente",
    "micro_tarea",
    "marca",
    "completada",
    "tiempo_estimado_micro_tarea_min",
    "tiempo_real_micro_tarea_min",
    "desviacion_micro_tarea_min",
    "hora_inicio_micro_tarea",
    "hora_fin_micro_tarea",
    "foto_evidencia_url",
]

RUTAS_HEADERS = [
    "fecha",
    "anio",
    "mes",
    "semana_iso",
    "dia_semana",
    "estado_ruta",
    "supervisor",
    "reponedor",
    "pdvs_planificados",
    "pdvs_completados",
    "pdvs_en_progreso",
    "pdvs_en_traslado",
    "pdvs_pendientes",
    "cobertura_pct",
    "distancia_total_planificada_km",
    "tiempo_total_planificado_min",
    "tiempo_total_real_min",
    "tiempo_traslado_real_min",
    "tiempo_pdv_real_min",
    "micro_tareas_total",
    "micro_tareas_completadas",
    "evidencias_fotos_total",
]

INCIDENCIAS_HEADERS = [
    "fecha",
    "hora_reporte",
    "supervisor",
    "reponedor",
    "pdv_codigo",
    "pdv_nombre",
    "mercado",
    "tipo_cliente",
    "categoria",
    "severidad",
    "estado",
    "afecta_entrega",
    "cantidad_afectada",
    "descripcion",
    "accion_tomada",
    "latitud",
    "longitud",
    "foto_url",
]


def _date_parts(value: date) -> dict:
    iso = value.isocalendar()
    return {
        "fecha": value.isoformat(),
        "anio": value.year,
        "mes": value.month,
        "semana_iso": iso.week,
        "dia_semana": DIAS_ES[value.weekday()],
    }


def _local_dt(value: datetime | None) -> str:
    if not value:
        return ""
    return value.astimezone(LOCAL_TZ).strftime("%Y-%m-%d %H:%M:%S")


def _point_lat_lng(value) -> tuple[float | str, float | str]:
    if value is None:
        return "", ""
    shape = to_shape(value)
    return round(shape.y, 8), round(shape.x, 8)


def _absolute_url(request: Request, path: str | None) -> str:
    if not path:
        return ""
    if path.startswith(("http://", "https://", "data:", "blob:")):
        return path
    return f"{str(request.base_url).rstrip('/')}/{path.lstrip('/')}"


def _pct(numerator: int | float, denominator: int | float) -> float:
    return round((numerator / denominator) * 100, 2) if denominator else 0.0


def _visitas_query(db: Session, fecha_desde: date | None, fecha_hasta: date | None):
    query = (
        db.query(Visita)
        .options(
            joinedload(Visita.pdv),
            joinedload(Visita.ruta).joinedload(Ruta.reponedor),
            joinedload(Visita.ejecuciones).joinedload(EjecucionMicroTarea.micro_tarea),
        )
        .join(Ruta)
    )
    if fecha_desde:
        query = query.filter(Ruta.fecha >= fecha_desde)
    if fecha_hasta:
        query = query.filter(Ruta.fecha <= fecha_hasta)
    return query.order_by(Ruta.fecha.desc(), Ruta.reponedor_id, Visita.orden_planificado).all()


def _rutas_query(db: Session, fecha_desde: date | None, fecha_hasta: date | None):
    query = db.query(Ruta).options(
        joinedload(Ruta.reponedor),
        joinedload(Ruta.visitas).joinedload(Visita.pdv),
        joinedload(Ruta.visitas).joinedload(Visita.ejecuciones).joinedload(EjecucionMicroTarea.micro_tarea),
    )
    if fecha_desde:
        query = query.filter(Ruta.fecha >= fecha_desde)
    if fecha_hasta:
        query = query.filter(Ruta.fecha <= fecha_hasta)
    return query.order_by(Ruta.fecha.desc(), Ruta.reponedor_id).all()


def _incidencias_query(db: Session, fecha_desde: date | None, fecha_hasta: date | None):
    query = (
        db.query(ReporteIncidencia)
        .options(
            joinedload(ReporteIncidencia.visita).joinedload(Visita.ruta).joinedload(Ruta.reponedor),
            joinedload(ReporteIncidencia.pdv),
            joinedload(ReporteIncidencia.reponedor),
        )
        .join(Visita, Visita.id == ReporteIncidencia.visita_id)
        .join(Ruta, Ruta.id == Visita.ruta_id)
    )
    if fecha_desde:
        query = query.filter(Ruta.fecha >= fecha_desde)
    if fecha_hasta:
        query = query.filter(Ruta.fecha <= fecha_hasta)
    return query.order_by(Ruta.fecha.desc(), ReporteIncidencia.created_at.desc()).all()


def build_visitas_rows(visitas: list[Visita], request: Request) -> list[dict]:
    rows = []
    for visita in visitas:
        date_cols = _date_parts(visita.ruta.fecha)
        checkin_lat, checkin_lng = _point_lat_lng(visita.coordenada_checkin)
        checkout_lat, checkout_lng = _point_lat_lng(visita.coordenada_checkout)
        completadas = sum(1 for e in visita.ejecuciones if e.completada)
        total = len(visita.ejecuciones)
        fotos = [_absolute_url(request, e.foto_evidencia_url) for e in visita.ejecuciones if e.foto_evidencia_url]
        desviacion_min = (
            visita.tiempo_ejecucion_min - visita.pdv.tiempo_visita_estimado_min
            if visita.tiempo_ejecucion_min is not None
            else ""
        )
        desviacion_pct = (
            round((desviacion_min / visita.pdv.tiempo_visita_estimado_min) * 100, 2)
            if desviacion_min != "" and visita.pdv.tiempo_visita_estimado_min
            else ""
        )
        rows.append(
            {
                **date_cols,
                "ruta_id": str(visita.ruta_id),
                "visita_id": str(visita.id),
                "orden_visita": visita.orden_planificado,
                "estado_visita": visita.estado.value,
                "supervisor": visita.ruta.reponedor.supervisor,
                "reponedor": visita.ruta.reponedor.nombre,
                "pdv_id": str(visita.pdv_id),
                "pdv_codigo": visita.pdv.codigo,
                "pdv_nombre": visita.pdv.nombre,
                "mercado": visita.pdv.mercado,
                "tipo_cliente": visita.pdv.tipo_cliente.value,
                "latitud_pdv": visita.pdv.latitud,
                "longitud_pdv": visita.pdv.longitud,
                "checkin_latitud": checkin_lat,
                "checkin_longitud": checkin_lng,
                "checkout_latitud": checkout_lat,
                "checkout_longitud": checkout_lng,
                "distancia_planificada_km": visita.distancia_desde_anterior_km or "",
                "distancia_ors_km": visita.ors_distancia_desde_anterior_km or "",
                "tiempo_traslado_planificado_min": visita.tiempo_traslado_desde_anterior_min or "",
                "tiempo_traslado_ors_min": visita.ors_tiempo_traslado_desde_anterior_min or "",
                "tiempo_traslado_real_min": visita.tiempo_traslado_real_min or "",
                "fuente_calculo_traslado": visita.traslado_fuente,
                "hora_inicio_traslado": _local_dt(visita.hora_inicio_traslado),
                "hora_fin_traslado": _local_dt(visita.hora_fin_traslado),
                "hora_inicio_visita": _local_dt(visita.hora_inicio_real),
                "hora_fin_visita": _local_dt(visita.hora_fin_real),
                "tiempo_pdv_estimado_min": visita.pdv.tiempo_visita_estimado_min,
                "tiempo_pdv_real_min": visita.tiempo_ejecucion_min or "",
                "desviacion_pdv_min": desviacion_min,
                "desviacion_pdv_pct": desviacion_pct,
                "desviacion_pdv_alerta": bool(
                    visita.tiempo_ejecucion_min
                    and visita.pdv.tiempo_visita_estimado_min
                    and visita.tiempo_ejecucion_min > visita.pdv.tiempo_visita_estimado_min * 1.5
                ),
                "micro_tareas_total": total,
                "micro_tareas_completadas": completadas,
                "micro_tareas_pendientes": total - completadas,
                "evidencias_fotos_total": len(fotos),
                "evidencias_fotos_urls": " | ".join(fotos),
            }
        )
    return rows


def build_micro_tareas_rows(visitas: list[Visita], request: Request) -> list[dict]:
    rows = []
    for visita in visitas:
        date_cols = _date_parts(visita.ruta.fecha)
        for ejecucion in visita.ejecuciones:
            tarea = ejecucion.micro_tarea
            desviacion = (
                ejecucion.tiempo_real_min - tarea.tiempo_estimado_minutos
                if ejecucion.tiempo_real_min is not None and tarea
                else ""
            )
            rows.append(
                {
                    **date_cols,
                    "ruta_id": str(visita.ruta_id),
                    "visita_id": str(visita.id),
                    "orden_visita": visita.orden_planificado,
                    "estado_visita": visita.estado.value,
                    "supervisor": visita.ruta.reponedor.supervisor,
                    "reponedor": visita.ruta.reponedor.nombre,
                    "pdv_codigo": visita.pdv.codigo,
                    "pdv_nombre": visita.pdv.nombre,
                    "mercado": visita.pdv.mercado,
                    "tipo_cliente": visita.pdv.tipo_cliente.value,
                    "micro_tarea_id": str(ejecucion.micro_tarea_id),
                    "micro_tarea": tarea.nombre if tarea else "",
                    "marca": tarea.marca if tarea else "",
                    "completada": ejecucion.completada,
                    "tiempo_estimado_micro_tarea_min": tarea.tiempo_estimado_minutos if tarea else "",
                    "tiempo_real_micro_tarea_min": ejecucion.tiempo_real_min or "",
                    "desviacion_micro_tarea_min": desviacion,
                    "hora_inicio_micro_tarea": _local_dt(ejecucion.hora_inicio),
                    "hora_fin_micro_tarea": _local_dt(ejecucion.hora_fin),
                    "foto_evidencia_url": _absolute_url(request, ejecucion.foto_evidencia_url),
                }
            )
    return rows


def build_rutas_rows(rutas: list[Ruta]) -> list[dict]:
    rows = []
    for ruta in rutas:
        visitas = list(ruta.visitas)
        total = len(visitas)
        completadas = sum(1 for v in visitas if v.estado == VisitaEstado.COMPLETADA)
        en_progreso = sum(1 for v in visitas if v.estado == VisitaEstado.EN_PROGRESO)
        en_traslado = sum(1 for v in visitas if v.estado == VisitaEstado.EN_TRASLADO)
        pendientes = sum(1 for v in visitas if v.estado == VisitaEstado.PENDIENTE)
        ejecuciones = [ej for visita in visitas for ej in visita.ejecuciones]
        rows.append(
            {
                **_date_parts(ruta.fecha),
                "ruta_id": str(ruta.id),
                "estado_ruta": ruta.estado.value,
                "supervisor": ruta.reponedor.supervisor,
                "reponedor": ruta.reponedor.nombre,
                "pdvs_planificados": total,
                "pdvs_completados": completadas,
                "pdvs_en_progreso": en_progreso,
                "pdvs_en_traslado": en_traslado,
                "pdvs_pendientes": pendientes,
                "cobertura_pct": _pct(completadas, total),
                "distancia_total_planificada_km": ruta.distancia_total_km,
                "tiempo_total_planificado_min": ruta.tiempo_total_estimado_min,
                "tiempo_total_real_min": ruta.tiempo_total_real_min or "",
                "tiempo_traslado_real_min": sum(v.tiempo_traslado_real_min or 0 for v in visitas),
                "tiempo_pdv_real_min": sum(v.tiempo_ejecucion_min or 0 for v in visitas),
                "micro_tareas_total": len(ejecuciones),
                "micro_tareas_completadas": sum(1 for ej in ejecuciones if ej.completada),
                "evidencias_fotos_total": sum(1 for ej in ejecuciones if ej.foto_evidencia_url),
            }
        )
    return rows


def build_incidencias_rows(incidencias: list[ReporteIncidencia], request: Request) -> list[dict]:
    rows = []
    for item in incidencias:
        rows.append(
            {
                "fecha": item.visita.ruta.fecha.isoformat(),
                "hora_reporte": _local_dt(item.created_at),
                "supervisor": item.visita.ruta.reponedor.supervisor,
                "reponedor": item.reponedor.nombre if item.reponedor else item.visita.ruta.reponedor.nombre,
                "pdv_codigo": item.pdv.codigo,
                "pdv_nombre": item.pdv.nombre,
                "mercado": item.pdv.mercado,
                "tipo_cliente": item.pdv.tipo_cliente.value,
                "categoria": item.categoria,
                "severidad": item.severidad,
                "estado": item.estado,
                "afecta_entrega": "SI" if item.afecta_entrega else "NO",
                "cantidad_afectada": item.cantidad_afectada,
                "descripcion": item.descripcion,
                "accion_tomada": item.accion_tomada,
                "latitud": item.latitud or "",
                "longitud": item.longitud or "",
                "foto_url": _absolute_url(request, item.foto_url),
            }
        )
    return rows


def _group_counts(rows: list[dict], key: str) -> list[dict]:
    grouped: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key) or "SIN_DATO")
        grouped[value] = grouped.get(value, 0) + 1
    ordered = sorted(grouped.items(), key=lambda item: item[1], reverse=True)
    total = sum(value for _, value in ordered) or 1
    cumulative = 0
    result = []
    for name, count in ordered:
        cumulative += count
        result.append({"nombre": name, "total": count, "acumulado_pct": round((cumulative / total) * 100, 2)})
    return result


def _style_header(row) -> None:
    fill = PatternFill("solid", fgColor="C8102E")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_table(ws, title: str, headers: list[str], rows: list[dict], table_name: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="0F172A")
    ws.append(headers)
    _style_header(ws[2])
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    if rows:
        ref = f"A2:{ws.cell(row=ws.max_row, column=len(headers)).coordinate}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    ws.freeze_panes = "A3"
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col[:80])
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 34)


def _add_pareto_sheet(wb: Workbook, name: str, rows: list[dict], title: str) -> None:
    ws = wb.create_sheet(name)
    safe_name = name.lower().replace(" ", "_").replace("-", "_")
    _write_table(ws, title, ["nombre", "total", "acumulado_pct"], rows, f"tbl_{safe_name}")
    if not rows:
        return
    max_row = ws.max_row
    bar = BarChart()
    bar.title = title
    bar.y_axis.title = "Cantidad"
    bar.x_axis.title = "Categoria"
    bar.add_data(Reference(ws, min_col=2, min_row=2, max_row=max_row), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=3, max_row=max_row))
    bar.height = 9
    bar.width = 18

    line = LineChart()
    line.add_data(Reference(ws, min_col=3, min_row=2, max_row=max_row), titles_from_data=True)
    line.y_axis.axId = 200
    line.y_axis.title = "Acumulado %"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = 100
    line.dataLabels = DataLabelList()
    line.dataLabels.showVal = False
    bar += line
    ws.add_chart(bar, "E3")


def _xlsx_incidencias_response(rows: list[dict], filename: str) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    total = len(rows)
    abiertas = sum(1 for row in rows if row.get("estado") == "ABIERTO")
    altas = sum(1 for row in rows if row.get("severidad") == "ALTA")
    afectan = sum(1 for row in rows if row.get("afecta_entrega") == "SI")

    ws["A1"] = "Dashboard de incidencias - Industrias Venado"
    ws["A1"].font = Font(size=18, bold=True, color="0F172A")
    kpis = [
        ("Total reportes", total),
        ("Abiertos", abiertas),
        ("Severidad alta", altas),
        ("Afectan entrega", afectan),
    ]
    for index, (label, value) in enumerate(kpis, start=1):
        col = 1 + (index - 1) * 3
        ws.cell(row=3, column=col, value=label)
        ws.cell(row=4, column=col, value=value)
        ws.cell(row=3, column=col).font = Font(bold=True, color="64748B")
        ws.cell(row=4, column=col).font = Font(size=20, bold=True, color="C8102E" if index in (2, 3, 4) else "174EA6")

    categoria = _group_counts(rows, "categoria")
    mercado = _group_counts(rows, "mercado")
    severidad = _group_counts(rows, "severidad")
    for title, data, start_col in [
        ("Pareto categoria", categoria[:8], 1),
        ("Top mercados", mercado[:8], 5),
        ("Severidad", severidad[:5], 9),
    ]:
        ws.cell(row=7, column=start_col, value=title).font = Font(bold=True, color="0F172A")
        ws.cell(row=8, column=start_col, value="nombre")
        ws.cell(row=8, column=start_col + 1, value="total")
        _style_header([ws.cell(row=8, column=start_col), ws.cell(row=8, column=start_col + 1)])
        for offset, item in enumerate(data, start=9):
            ws.cell(row=offset, column=start_col, value=item["nombre"])
            ws.cell(row=offset, column=start_col + 1, value=item["total"])
        if data:
            chart = BarChart()
            chart.title = title
            chart.add_data(Reference(ws, min_col=start_col + 1, min_row=8, max_row=8 + len(data)), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=start_col, min_row=9, max_row=8 + len(data)))
            chart.height = 7
            chart.width = 11
            ws.add_chart(chart, ws.cell(row=17, column=start_col).coordinate)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 12

    detail = wb.create_sheet("Incidencias")
    _write_table(detail, "Detalle de incidencias", INCIDENCIAS_HEADERS, rows, "tbl_incidencias")
    _add_pareto_sheet(wb, "Pareto Categoria", categoria, "Pareto de defectos por categoria")
    _add_pareto_sheet(wb, "Pareto Mercado", mercado, "Pareto de defectos por mercado")
    _add_pareto_sheet(wb, "Severidad", severidad, "Distribucion por severidad")

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _csv_value(value) -> str | int | float:
    if isinstance(value, bool):
        return "SI" if value else "NO"
    if value is None:
        return ""
    return value


def _csv_response(rows: list[dict], headers: list[str], filename: str) -> StreamingResponse:
    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.DictWriter(output, fieldnames=headers, delimiter=",", lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    writer.writerows([{key: _csv_value(row.get(key, "")) for key in headers} for row in rows])
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/exportar-bi")
def exportar_bi(
    request: Request,
    tabla: TablaBI = Query(TablaBI.visitas, description="visitas, micro_tareas o rutas"),
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    db: Session = Depends(get_db),
    _user=Depends(require_role("supervisor")),
) -> StreamingResponse:
    if tabla == TablaBI.rutas:
        rows = build_rutas_rows(_rutas_query(db, fecha_desde, fecha_hasta))
        return _csv_response(rows, RUTAS_HEADERS, "venado_bi_rutas.csv")

    visitas = _visitas_query(db, fecha_desde, fecha_hasta)
    if tabla == TablaBI.micro_tareas:
        rows = build_micro_tareas_rows(visitas, request)
        return _csv_response(rows, MICRO_TAREAS_HEADERS, "venado_bi_micro_tareas.csv")

    rows = build_visitas_rows(visitas, request)
    return _csv_response(rows, VISITAS_HEADERS, "venado_bi_visitas.csv")


@router.get("/exportar-incidencias-excel")
def exportar_incidencias_excel(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    db: Session = Depends(get_db),
    _user=Depends(require_role("supervisor")),
) -> StreamingResponse:
    rows = build_incidencias_rows(_incidencias_query(db, fecha_desde, fecha_hasta), request)
    return _xlsx_incidencias_response(rows, "venado_incidencias_calidad.xlsx")


@router.get("/power-bi/visitas")
def power_bi_visitas(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    db: Session = Depends(get_db),
    _user=Depends(require_role("supervisor")),
) -> list[dict]:
    return build_visitas_rows(_visitas_query(db, fecha_desde, fecha_hasta), request)


@router.get("/power-bi/micro-tareas")
def power_bi_micro_tareas(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    db: Session = Depends(get_db),
    _user=Depends(require_role("supervisor")),
) -> list[dict]:
    return build_micro_tareas_rows(_visitas_query(db, fecha_desde, fecha_hasta), request)


@router.get("/power-bi/rutas")
def power_bi_rutas(
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    db: Session = Depends(get_db),
    _user=Depends(require_role("supervisor")),
) -> list[dict]:
    return build_rutas_rows(_rutas_query(db, fecha_desde, fecha_hasta))
